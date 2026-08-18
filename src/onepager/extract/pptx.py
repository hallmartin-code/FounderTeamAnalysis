"""PPTX extraction: recursive shape walk, tables, and speaker notes."""

from __future__ import annotations

import subprocess
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from ..config import soffice_path
from ..util.logging import get_logger
from .types import Deck, ExtractionError, SlideContent

_log = get_logger()

#: A slide yielding less than this many characters is treated as visually-authored.
_THIN_SLIDE_CHARS = 40


def _walk(shape, out: list[str], depth: int = 0) -> None:
    """Collect text from a shape, recursing into groups and tables."""
    if depth > 6:  # deeply nested groups are pathological; stop rather than recurse forever
        return
    try:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for child in shape.shapes:
                _walk(child, out, depth + 1)
            return
    except (AttributeError, NotImplementedError, ValueError):
        # Rotated/OLE/placeholder shapes can raise on shape_type; fall through to text.
        pass

    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            cells = [c.text.strip() for c in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                out.append(line)
        return

    if getattr(shape, "has_text_frame", False):
        for para in shape.text_frame.paragraphs:
            line = "".join(run.text for run in para.runs).strip()
            if not line:
                line = para.text.strip()
            if line:
                out.append(line)


def _slide_title(slide) -> str | None:
    try:
        if slide.shapes.title is not None:
            title = slide.shapes.title.text.strip()
            return title[:120] or None
    except (AttributeError, ValueError):
        pass
    return None


def _notes(slide) -> str | None:
    try:
        if slide.has_notes_slide:
            text = slide.notes_slide.notes_text_frame.text.strip()
            return text or None
    except (AttributeError, ValueError):
        return None
    return None


def extract_pptx_text(path: Path) -> Deck:
    deck = Deck(source_path=str(path), source_kind=path.suffix.lower().lstrip("."))
    try:
        prs = Presentation(str(path))
    except Exception as exc:
        raise ExtractionError(
            f"{path.name} could not be opened as a PowerPoint file "
            f"({type(exc).__name__}: {exc}). Legacy .ppt must be converted first."
        ) from exc

    for i, slide in enumerate(prs.slides, start=1):
        lines: list[str] = []
        for shape in slide.shapes:
            _walk(shape, lines)
        body = "\n".join(dict.fromkeys(lines))  # order-preserving dedupe of repeated runs
        # Many decks skip the title placeholder entirely; fall back to the first line.
        title = _slide_title(slide) or (lines[0][:120] if lines else None)
        deck.slides.append(
            SlideContent(index=i, title=title, text=body.strip(), speaker_notes=_notes(slide))
        )

    if not deck.slides:
        raise ExtractionError(f"{path.name} contains zero slides.")
    return deck


def convert_via_libreoffice(path: Path, outdir: Path) -> Path:
    """Convert any Office deck to PDF. Raises ExtractionError if LibreOffice is absent."""
    soffice = soffice_path()
    if not soffice:
        raise ExtractionError(
            "LibreOffice (soffice) was not found on PATH, so this file cannot be "
            "converted for image-based analysis. Install LibreOffice, or export the "
            "deck to PDF yourself and run onepager on the PDF."
        )
    cmd = [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(path)]
    try:
        proc = subprocess.run(cmd, capture_output=True, timeout=180, check=False)
    except subprocess.TimeoutExpired as exc:
        raise ExtractionError("LibreOffice conversion timed out after 180s.") from exc
    produced = outdir / (path.stem + ".pdf")
    if not produced.exists():
        detail = (proc.stderr or proc.stdout or b"").decode("utf-8", "replace").strip()
        raise ExtractionError(f"LibreOffice failed to convert {path.name}: {detail or 'no output'}")
    return produced


def pptx_to_deck(path: Path, vision: bool | None = None) -> Deck:
    """Text pass first; fall back to LibreOffice + rasterization for visual decks."""
    from .pdf import pdf_to_deck  # local import: avoids a cycle at module load

    if path.suffix.lower() == ".ppt":
        with tempfile.TemporaryDirectory() as tmp:
            converted = convert_via_libreoffice(path, Path(tmp))
            deck = pdf_to_deck(converted, vision=vision)
        deck.source_path = str(path)
        deck.source_kind = "ppt"
        deck.notes.append("legacy .ppt converted to PDF via LibreOffice")
        return deck

    deck = extract_pptx_text(path)
    thin = [s.index for s in deck.slides if s.char_count < _THIN_SLIDE_CHARS]
    mostly_visual = len(thin) > len(deck.slides) / 2

    if vision is False:
        return deck
    if vision is True or mostly_visual:
        if mostly_visual:
            deck.notes.append(
                f"{len(thin)}/{len(deck.slides)} slides yielded almost no shape text; "
                "converting to PDF for image analysis"
            )
        try:
            with tempfile.TemporaryDirectory() as tmp:
                converted = convert_via_libreoffice(path, Path(tmp))
                rendered = pdf_to_deck(converted, vision=True)
        except ExtractionError as exc:
            deck.notes.append(f"image fallback unavailable: {exc}")
            _log.warning("%s", exc)
            return deck
        # Keep the PPTX text and speaker notes; borrow the rendered page images.
        for slide, page in zip(deck.slides, rendered.slides, strict=False):
            slide.image_b64 = page.image_b64
            if not slide.text:
                slide.text = page.text
                slide.char_count = len(slide.text)
        deck.used_vision = True
    return deck
